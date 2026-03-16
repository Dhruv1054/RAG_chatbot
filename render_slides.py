"""
Render each PPTX slide as a PNG preview image using Pillow.
Reads shapes/text directly from the PPTX XML via python-pptx.
"""

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
import os, textwrap

SLIDE_W_PX = 1333
SLIDE_H_PX = 750
SCALE = SLIDE_W_PX / (13.33 * 914400)  # EMU → pixel

OUT_DIR = "/home/user/RAG_chatbot/slides_preview"
os.makedirs(OUT_DIR, exist_ok=True)

PPTX_PATH = "/home/user/RAG_chatbot/RAG_Chatbot_Presentation.pptx"

# ── font helpers ───────────────────────────────────────────────────────────────
def get_font(size_pt, bold=False):
    try:
        style = "Bold" if bold else "Regular"
        for name in [f"/usr/share/fonts/truetype/dejavu/DejaVuSans-{style}.ttf",
                     f"/usr/share/fonts/truetype/liberation/LiberationSans-{style}.ttf",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"]:
            if os.path.exists(name):
                return ImageFont.truetype(name, int(size_pt * 1.33))
    except Exception:
        pass
    return ImageFont.load_default()


def emu_to_px(emu):
    return int(emu * SCALE)


def rgb_from_pptx(color_obj):
    """Return (r,g,b) tuple or None."""
    try:
        c = color_obj.rgb
        return (c.red, c.green, c.blue)
    except Exception:
        return None


# ── render one slide ───────────────────────────────────────────────────────────
def render_slide(slide, idx):
    # background
    try:
        bg_fill = slide.background.fill
        bg_rgb = rgb_from_pptx(bg_fill.fore_color) or (13, 27, 42)
    except Exception:
        bg_rgb = (13, 27, 42)

    img = Image.new("RGB", (SLIDE_W_PX, SLIDE_H_PX), bg_rgb)
    draw = ImageDraw.Draw(img)

    for shape in slide.shapes:
        left   = emu_to_px(shape.left   or 0)
        top    = emu_to_px(shape.top    or 0)
        width  = emu_to_px(shape.width  or 0)
        height = emu_to_px(shape.height or 0)

        # ── filled rectangles ──────────────────────────────────────────────
        if shape.shape_type == 1:  # MSO_SHAPE_TYPE.AUTO_SHAPE (rectangle)
            try:
                fill = shape.fill
                if fill.type is not None:
                    fc = rgb_from_pptx(fill.fore_color)
                    if fc:
                        draw.rectangle([left, top, left+width, top+height], fill=fc)
                # border
                try:
                    lc = rgb_from_pptx(shape.line.color)
                    if lc and shape.line.width and shape.line.width > 0:
                        lw = max(1, int(shape.line.width / 12700))
                        draw.rectangle([left, top, left+width, top+height],
                                       outline=lc, width=lw)
                except Exception:
                    pass
            except Exception:
                pass

        # ── text boxes ────────────────────────────────────────────────────
        if shape.has_text_frame:
            tf = shape.text_frame
            y_cursor = top + 4
            for para in tf.paragraphs:
                line_parts = []
                max_pt = 12
                bold = False
                color = (255, 255, 255)
                align = "left"
                for run in para.runs:
                    t = run.text
                    if not t:
                        continue
                    pt = run.font.size / 12700 if run.font.size else 12
                    max_pt = max(max_pt, pt)
                    b = run.font.bold or False
                    bold = bold or b
                    c = rgb_from_pptx(run.font.color) or (255, 255, 255)
                    if c != (255,255,255) or not color:
                        color = c
                    line_parts.append(t)
                try:
                    al = para.alignment
                    if al and al.name == "CENTER":
                        align = "center"
                    elif al and al.name == "RIGHT":
                        align = "right"
                except Exception:
                    pass

                line_text = "".join(line_parts)
                if not line_text.strip():
                    y_cursor += int(max_pt * 1.33 * 0.5)
                    continue

                font = get_font(max_pt, bold)
                # wrap text
                chars_per_line = max(1, int(width / (max_pt * 0.75))) if width > 0 else 60
                wrapped = textwrap.wrap(line_text, width=chars_per_line) or [line_text]
                line_h = int(max_pt * 1.33 * 1.25)

                for wline in wrapped:
                    if y_cursor > top + height + 20:
                        break
                    tw = draw.textlength(wline, font=font) if hasattr(draw, 'textlength') else len(wline) * max_pt * 0.6
                    if align == "center":
                        tx = left + (width - tw) // 2
                    elif align == "right":
                        tx = left + width - tw - 4
                    else:
                        tx = left + 4

                    # subtle shadow
                    draw.text((tx+1, y_cursor+1), wline, font=font, fill=(0,0,0))
                    draw.text((tx, y_cursor), wline, font=font, fill=color)
                    y_cursor += line_h

    return img


# ── main ───────────────────────────────────────────────────────────────────────
prs = Presentation(PPTX_PATH)
saved = []
for i, slide in enumerate(prs.slides):
    img = render_slide(slide, i)
    path = os.path.join(OUT_DIR, f"slide_{i+1:02d}.png")
    img.save(path, "PNG")
    saved.append(path)
    print(f"  Saved: {path}")

print(f"\nDone — {len(saved)} slides rendered to {OUT_DIR}/")
